"""
Pengujian logika inti (tanpa Streamlit).

Jalankan dari folder project:
    python -m unittest discover -s tests -v
atau (jika pytest terpasang):
    pytest -v
"""

import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import numpy as np

from core.extractor import extract_text_from_txt, process_files_bytes, split_text_into_units
from core.pipeline import run_analysis
from core.preprocessing import build_vocabulary, preprocess_all, remove_punctuation
from core.similarity import (
    build_binary_matrix,
    build_similarity_matrix,
    cosine_from_counts,
    cosine_similarity_pair,
    interpret_similarity,
)

# Lima status pada PDF perhitungan manual
STATUS = [
    "Saya suka sayur dan buah",
    "Saya suka sayur tapi tak suka kimchi",
    "Hampir semua sayur saya suka. Pare pun doyan",
    "Saya suka sayur sayuran di Tesco.",
    "Saya suka makan sayur kangkung",
]

# Nilai pada PDF untuk pasangan yang TIDAK melibatkan S5 (semuanya sudah benar)
PDF_VALUES = {
    ("S1", "S2"): 0.548,
    ("S1", "S3"): 0.474,
    ("S1", "S4"): 0.548,
    ("S2", "S3"): 0.433,
    ("S2", "S4"): 0.500,
    ("S3", "S4"): 0.433,
}

# Pasangan dengan S5. Di PDF S5 ditulis berjumlah 6 kata (0.548, 0.500, 0.433, 0.500),
# padahal "saya suka makan sayur kangkung" hanya 5 kata → nilai yang benar:
CORRECT_S5_VALUES = {
    ("S1", "S5"): 0.600,  # 3 / (√5 × √5)
    ("S2", "S5"): 0.548,  # 3 / (√6 × √5)
    ("S3", "S5"): 0.474,  # 3 / (√8 × √5)
    ("S4", "S5"): 0.548,  # 3 / (√6 × √5)
}


class TestPreprocessing(unittest.TestCase):
    def test_remove_punctuation(self):
        self.assertEqual(remove_punctuation("saya suka sayur. pare, pun!"), "saya suka sayur pare pun")
        self.assertEqual(remove_punctuation("“halo” — dunia_ok"), "halo dunia ok")

    def test_tokens_and_unique_counts_match_pdf(self):
        res = preprocess_all(STATUS)
        self.assertEqual(res[1]["tokens"].count("suka"), 2)  # "suka" muncul dua kali di S2
        self.assertEqual([len(set(r["tokens"])) for r in res], [5, 6, 8, 6, 5])

    def test_vocabulary_order_matches_pdf_table(self):
        vocab = build_vocabulary(preprocess_all(STATUS))
        self.assertEqual(
            vocab,
            ["saya", "suka", "sayur", "dan", "buah", "tapi", "tak", "kimchi", "hampir",
             "semua", "pare", "pun", "doyan", "sayuran", "di", "tesco", "makan", "kangkung"],
        )


class TestSimilarity(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.res = run_analysis(STATUS)
        cls.values = {
            (r["Status 1"], r["Status 2"]): r["Cosine Similarity"] for _, r in cls.res["pairs_df"].iterrows()
        }

    def test_binary_totals(self):
        self.assertEqual(self.res["binary_total"].tolist(), [5, 6, 8, 6, 5])

    def test_common_words(self):
        self.assertEqual(self.res["common_words"], ["saya", "suka", "sayur"])

    def test_values_match_pdf(self):
        for pair, expected in PDF_VALUES.items():
            self.assertAlmostEqual(self.values[pair], expected, places=3, msg=str(pair))

    def test_s5_pairs_use_five_words(self):
        for pair, expected in CORRECT_S5_VALUES.items():
            self.assertAlmostEqual(self.values[pair], expected, places=3, msg=str(pair))

    def test_ten_pairs_and_details(self):
        self.assertEqual(len(self.res["pairs_df"]), 10)
        self.assertEqual(len(self.res["details"]), 10)
        d = self.res["details"][0]
        self.assertEqual((d["id_a"], d["id_b"], d["count_a"], d["count_b"], d["shared_count"]), ("S1", "S2", 5, 6, 3))

    def test_matrix_symmetric_diagonal_one(self):
        m = self.res["sim_matrix"]
        np.testing.assert_allclose(m, m.T)
        np.testing.assert_allclose(np.diag(m), 1.0)

    def test_vector_cosine_equals_shortcut_formula(self):
        """cos(x,y) dari vektor binary == kata_sama / (√a × √b) untuk semua pasangan."""
        for d in self.res["details"]:
            self.assertAlmostEqual(d["similarity"], self.values[(d["id_a"], d["id_b"])])
            self.assertAlmostEqual(
                d["similarity"], cosine_from_counts(d["shared_count"], d["count_a"], d["count_b"])
            )

    def test_zero_vector_and_no_overlap(self):
        self.assertEqual(cosine_similarity_pair(np.zeros(3), np.array([1.0, 0, 1])), 0.0)
        res = run_analysis(["apel jeruk", "mobil motor"])
        self.assertEqual(res["pairs_df"].loc[0, "Cosine Similarity"], 0.0)

    def test_identical_texts(self):
        res = run_analysis(["Saya suka sayur", "saya suka sayur!"])
        self.assertAlmostEqual(res["pairs_df"].loc[0, "Cosine Similarity"], 1.0)

    def test_interpretation_boundaries(self):
        self.assertEqual(interpret_similarity(0.6), "Tinggi")
        self.assertEqual(interpret_similarity(0.5999999999999999), "Tinggi")  # galat floating point
        self.assertEqual(interpret_similarity(0.548), "Sedang")
        self.assertEqual(interpret_similarity(0.8), "Sangat Tinggi")
        self.assertEqual(interpret_similarity(0.0), "Sangat Rendah")
        self.assertEqual(self.res["pairs_df"].set_index(["Status 1", "Status 2"]).loc[("S1", "S5"), "Interpretasi"], "Tinggi")


class TestPipeline(unittest.TestCase):
    def test_needs_two_texts(self):
        with self.assertRaises(ValueError):
            run_analysis(["hanya satu", "   "])

    def test_labels_keep_original_numbers_when_blank_skipped(self):
        res = run_analysis(["a b", "", "a c"], labels=["S1", "S2", "S3"])
        self.assertEqual(res["status_ids"], ["S1", "S3"])

    def test_sources_column(self):
        res = run_analysis(["a b", "a c"], sources=["x.txt", "y.txt"])
        self.assertIn("Sumber", res["prep_df"].columns)

    def test_text_without_tokens(self):
        res = run_analysis(["!!!", "saya suka"])
        self.assertEqual(res["pairs_df"].loc[0, "Cosine Similarity"], 0.0)


class TestExtractor(unittest.TestCase):
    def test_txt_and_split_modes(self):
        raw = "baris satu\n\nbaris dua\n".encode("utf-8")
        text, warn = extract_text_from_txt(raw)
        self.assertEqual(warn, "")
        self.assertEqual(len(split_text_into_units(text, "file")), 1)
        self.assertEqual(split_text_into_units(text, "line"), ["baris satu", "baris dua"])

    def test_process_files_bytes(self):
        data = "Saya suka sayur\nSaya suka buah\n".encode("utf-8")
        out = process_files_bytes([("a.txt", data), ("x.xyz", b"?")], "line")
        self.assertEqual(len(out[0]["units"]), 2)
        self.assertTrue(out[1]["warning"])

    def test_docx_roundtrip(self):
        try:
            from docx import Document
        except ImportError:
            self.skipTest("python-docx belum terpasang")
        doc = Document()
        doc.add_paragraph("Saya suka sayur")
        doc.add_paragraph("Saya suka buah")
        buf = io.BytesIO()
        doc.save(buf)
        out = process_files_bytes([("a.docx", buf.getvalue())], "line")
        self.assertEqual(out[0]["units"], ["Saya suka sayur", "Saya suka buah"])


if __name__ == "__main__":
    unittest.main()
